import json
import hashlib
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy import delete, func, select

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.core.db import SessionLocal
from app.models.control import Control
from app.models.control_audit_log import ControlAuditLog
from app.models.user import User


SOURCE_SHEET = "Model raportare"
HEADER_ROW = 4
DATA_START_ROW = 6
EXPECTED_COLUMNS = 15
IMPORT_CONFIRM = os.getenv("GFN_REAL_IMPORT_CONFIRM", "").lower() == "true"
IMPORT_REPLACE = os.getenv("GFN_REAL_IMPORT_REPLACE", "").lower() == "true"
IMPORT_FILE = Path(os.getenv("GFN_REAL_IMPORT_FILE", "/app/imports/centralizare_controale_reale.xlsx"))
IMPORT_BATCH = os.getenv("GFN_REAL_IMPORT_BATCH", f"real_excel_{datetime.now():%Y%m%d%H%M%S}")
REPORT_FILE = Path(os.getenv("GFN_REAL_IMPORT_REPORT", f"/tmp/{IMPORT_BATCH}_report.json"))

CANONICAL_GUARDS = {
    "brasov": "Garda Forestiera Brasov",
    "bucuresti": "Garda Forestiera Bucuresti",
    "cluj": "Garda Forestiera Cluj",
    "focsani": "Garda Forestiera Focsani",
    "oradea": "Garda Forestiera Oradea",
    "ploiesti": "Garda Forestiera Ploiesti",
    "ramnicu valcea": "Garda Forestiera Ramnicu Valcea",
    "suceava": "Garda Forestiera Suceava",
    "timisoara": "Garda Forestiera Timisoara",
}

GUARD_SPECIAL_ALIASES = {
    "0radea": "oradea",
    "sv": "suceava",
    "gf sv": "suceava",
    "bt": "suceava",
    "gfbt": "suceava",
    "gf mh": "ramnicu valcea",
    "mh": "ramnicu valcea",
    "mehedinti": "ramnicu valcea",
    "gfj mehedinti": "ramnicu valcea",
    "gfj gorj": "ramnicu valcea",
    "garda forestiera judeteana gorj": "ramnicu valcea",
    "gorj": "ramnicu valcea",
    "gfj olt": "ramnicu valcea",
    "dolj": "ramnicu valcea",
    "gfj bacau": "focsani",
}

ROMANIA_BOUNDS = (43.0, 49.0, 20.0, 30.5)
INSTITUTION_WORDS = {
    "ipj", "politie", "jandarmi", "personal", "echipa", "reprezentanti",
    "reprezentant", "comisariat", "garda", "ocol", "silvic", "silivc",
    "gf", "gfj", "ds", "gnm", "isu", "aba", "dsv", "sop", "sjpt",
}

PERSON_TITLE_RE = (
    r"(?:ing(?:iner)?|insp(?:ector)?|ins|cons(?:ilier)?|cns|"
    r"scms|cms|comisar|ag(?:ent)?|sef|adj|af|ap|ppc|pr|p[aă]d(?:urar)?|tehn(?:ician)?)"
)
PARTNER_TITLE_RE = r"(?:scms|cms|comisar|ag(?:ent)?|sef|adj|af|ap|ppc)"
AFFILIATION_RE = (
    r"(?:\bdin\s+cadrul\b|\bimpreuna\s+cu\b|\breprezentant(?:i|ii)?\b|"
    r"\blucrator(?:i)?\b|\bpersonal\b|\bagenti?\s+sop\b|\bactiune\b|"
    r"\bsef\s+district\b|\bconsilier\b|\bpadurar(?:i)?\b|"
    r"\s+-\s*(?:gfj?|g\.?f\.?j?\.?|ipj|i\.?p\.?j\.?|gnm|isu|aba|dsv|os)\b)"
)
NAME_STOP_WORDS = INSTITUTION_WORDS | {
    "agenti", "agent", "control", "controlului", "constatare", "constatari",
    "dispus", "dispusa", "dispuse", "lucrari", "materializarea", "teren",
    "data", "anului", "cadrul", "organe", "impreuna", "coordonarea",
    "parchetului", "judecatoria", "reprezentantii", "actiune", "sef",
    "district", "delegat", "politia", "politiei", "consilier", "superior",
    "silvicultura", "primaria", "padurari", "lucrator", "lucratori", "vamali", "os",
    "gf", "gfj", "ds", "gnm", "isu", "aba", "dsv", "sop", "sjpt",
    "suceva", "brasov", "mures", "iasi", "bacau", "bc", "is",
}


@dataclass
class ParsedRow:
    excel_row: int
    source_number: str
    date_start: date
    date_end: date
    control_type: str
    result: str
    payload: dict[str, Any]


def clean_spaces(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\u00a0", " ")).strip()


def ascii_key(value: Any) -> str:
    text = clean_spaces(value).lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def normalize_guard(value: Any) -> Optional[str]:
    key = ascii_key(value).replace("0radea", "oradea")
    if not key:
        return None
    if key in GUARD_SPECIAL_ALIASES:
        return CANONICAL_GUARDS[GUARD_SPECIAL_ALIASES[key]]
    if "brasov" in key:
        return CANONICAL_GUARDS["brasov"]
    if "bucuresti" in key:
        return CANONICAL_GUARDS["bucuresti"]
    if "cluj" in key:
        return CANONICAL_GUARDS["cluj"]
    if "focsani" in key:
        return CANONICAL_GUARDS["focsani"]
    if "oradea" in key:
        return CANONICAL_GUARDS["oradea"]
    if "ploiesti" in key:
        return CANONICAL_GUARDS["ploiesti"]
    if "valcea" in key or key in {"rm valcea", "rmv"}:
        return CANONICAL_GUARDS["ramnicu valcea"]
    if "suceava" in key:
        return CANONICAL_GUARDS["suceava"]
    if "timisoara" in key or "timisora" in key:
        return CANONICAL_GUARDS["timisoara"]
    return None


def normalize_judet(value: Any) -> str:
    key = ascii_key(value)
    if not key:
        return ""
    aliases = {
        "bistrita nasaud": "Bistrita-Nasaud",
        "caras severin": "Caras-Severin",
        "satu mare": "Satu Mare",
        "valcea": "Valcea",
        "bucuresti": "Bucuresti",
    }
    return aliases.get(key, " ".join(part.capitalize() for part in key.split()))


def normalize_display_text(value: Any) -> str:
    return clean_spaces(value)


def make_date(day: int, month: int, year: int) -> Optional[date]:
    if year == 206:
        year = 2026
    elif 100 <= year <= 999:
        return None
    if year < 100:
        year += 2000
    if month > 12 and 1 <= day <= 12:
        day, month = month, day
    if year < 2000 or year > 2100:
        return None
    try:
        return date(year, month, day)
    except ValueError:
        return None


def parse_control_dates(value: Any) -> tuple[Optional[date], Optional[date], Optional[str]]:
    if isinstance(value, datetime):
        parsed = value.date()
        return (parsed, parsed, None) if 2000 <= parsed.year <= 2100 else (None, None, "out_of_range")
    if isinstance(value, date):
        return (value, value, None) if 2000 <= value.year <= 2100 else (None, None, "out_of_range")
    if isinstance(value, (int, float)) and 20_000 <= float(value) <= 80_000:
        converted = from_excel(value)
        parsed = converted.date() if isinstance(converted, datetime) else converted
        return (
            (parsed, parsed, "excel_serial")
            if 2000 <= parsed.year <= 2100
            else (None, None, "out_of_range")
        )

    raw = clean_spaces(value)
    if not raw:
        return None, None, "missing"

    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00")).date()
        if 2000 <= parsed.year <= 2100:
            return parsed, parsed, None
        return None, None, "out_of_range"
    except ValueError:
        pass

    text = raw.lower().replace("..", ".").replace("/", ".").replace(",", ".")
    text = re.sub(r"\s+", " ", text).strip()
    month_words = {
        "ianuarie": "01", "ian": "01",
        "februarie": "02", "feb": "02",
        "martie": "03", "mar": "03",
        "aprilie": "04", "apr": "04",
        "mai": "05",
        "iunie": "06", "iun": "06",
        "iulie": "07", "iul": "07",
        "august": "08", "aug": "08",
        "septembrie": "09", "sept": "09", "sep": "09",
        "octombrie": "10", "oct": "10",
        "noiembrie": "11", "nov": "11",
        "decembrie": "12", "dec": "12",
    }
    for name, month in month_words.items():
        text = re.sub(rf"(?i){name}", f".{month}.", text)
    text = re.sub(r"\.{2,}", ".", text)
    # Cautam mai intai o data completa. In expresii precum 12-19.02.2026,
    # prima cifra este inceputul intervalului, nu anul/luna unei alte date.
    full_matches = list(
        re.finditer(
            r"(?<!\d)(\d{1,2})\s*[.\-]\s*(\d{1,2})\s*[.\-]\s*(\d{3,4})(?!\d)",
            text,
        )
    )
    full_dates = [
        make_date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
        for match in full_matches
    ]
    full_dates = [item for item in full_dates if item]
    if len(full_dates) >= 2:
        return min(full_dates), max(full_dates), "interval_full"

    if len(full_dates) == 1:
        end = full_dates[0]
        prefix = text[:full_matches[0].start()].strip(" .,-")
        day_month = re.search(r"(\d{1,2})\s*[.\-]\s*(\d{1,2})\s*$", prefix)
        day_only = re.search(r"(\d{1,2})\s*$", prefix)
        if day_month:
            start = make_date(int(day_month.group(1)), int(day_month.group(2)), end.year)
            if start:
                return min(start, end), max(start, end), "interval_inferred_year"
        if day_only and re.search(r"[-,]", text[:full_matches[0].start()]):
            start = make_date(int(day_only.group(1)), end.month, end.year)
            if start:
                return min(start, end), max(start, end), "interval_inferred_month"
        return end, end, None

    return None, None, "unparsed"


def parse_coord(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = clean_spaces(value).replace(",", ".")
    match = re.fullmatch(r"([+-]?\d{1,3}(?:\.\d+)?)\s*°?", text)
    if not match:
        return None
    try:
        return float(match.group(1))
    except ValueError:
        return None


def normalize_coordinates(lat_raw: Any, lon_raw: Any) -> tuple[Optional[dict[str, float]], str]:
    lat = parse_coord(lat_raw)
    lon = parse_coord(lon_raw)
    if lat is None or lon is None:
        return None, "missing_or_ambiguous"
    min_lat, max_lat, min_lon, max_lon = ROMANIA_BOUNDS
    if min_lat <= lat <= max_lat and min_lon <= lon <= max_lon:
        return {"lat": round(lat, 7), "lon": round(lon, 7)}, "valid"
    if min_lat <= lon <= max_lat and min_lon <= lat <= max_lon:
        return {"lat": round(lon, 7), "lon": round(lat, 7)}, "swapped"
    return None, "outside_romania"


def normalize_control_type(value: Any) -> str:
    key = ascii_key(value)
    if "sesiz" in key or "seziz" in key or "petit" in key:
        return "sesizare"
    if "operativ" in key:
        return "operativ"
    if "tematic" in key:
        return "tematic"
    if "rutin" in key:
        return "rutina"
    if key in {"fond", "control fond", "control de fond"}:
        return "fond"
    if "solicitar" in key or key in {"adresa", "notificare"}:
        return "solicitare"
    return "necunoscut"


def normalize_result(value: Any) -> tuple[str, str]:
    key = ascii_key(value)
    if not key or key in {"-", "nu este cazul", "nu a fost cazul"}:
        return "necunoscut", "necunoscut"
    if any(token in key for token in ("in lucru", "in curs", "desfasurare", "nefinalizat", "finalizare")):
        return "necunoscut", "necunoscut"
    has_conform = "conform" in key and "neconform" not in key
    has_nonconform = "neconform" in key or re.search(r"\bnc\b", key) is not None
    if has_conform and has_nonconform:
        return "mixt", "mixt"
    if "penal" in key or "parchet" in key:
        return "sesizare_penala", "neconform_sanctiune"
    if "avert" in key:
        return "avertisment", "neconform_sanctiune"
    if any(token in key for token in ("sanct", "amenda", "contrav", "confisc")):
        return "sanctiune", "neconform_sanctiune"
    if has_nonconform:
        return "neconform", "neconform_sanctiune"
    if has_conform:
        return "conform", "conform"
    return "necunoscut", "necunoscut"


def normalize_person_name(value: Any) -> str:
    text = clean_spaces(value).strip(" ,;+/-")
    if not text:
        return ""
    text = re.sub(
        rf"^(?:(?:{PERSON_TITLE_RE}|dr)\b\.?\s*)+",
        "",
        text,
        flags=re.IGNORECASE,
    )
    return " ".join(part.upper() if len(part) <= 2 and part.endswith(".") else part.capitalize() for part in text.split())


def split_raw_inspector_segments(value: Any) -> list[str]:
    if value is None:
        return []
    text = str(value).replace("\u00a0", " ").replace("\r", "\n")
    text = re.sub(
        rf"\s+(?=(?:{PERSON_TITLE_RE})\b\.?\s*)",
        " | ",
        text,
        flags=re.IGNORECASE,
    )
    parts = re.split(
        r"(?:\n+|\s{3,}|[,;:+/]+|\s+\b(?:si|și)\b\s+|\s+\|\s+)",
        text,
        flags=re.IGNORECASE,
    )
    return [part.strip() for part in parts if part and part.strip()]


def clean_person_candidate(value: Any) -> tuple[str, bool]:
    raw = clean_spaces(value).strip(" |,;:+/")
    if not raw:
        return "", False

    starts_as_partner = bool(
        re.match(rf"^(?:{PARTNER_TITLE_RE})\b\.?\s+", raw, flags=re.IGNORECASE)
    )
    explicit_partner_affiliation = bool(
        re.search(
            r"\s*-\s*(?:ipj|i\.?p\.?j\.?|gnm|isu|aba|dsv|politia)\b",
            raw,
            flags=re.IGNORECASE,
        )
    )
    raw = re.sub(
        rf"^(?:(?:{PERSON_TITLE_RE}|dr)\b\.?\s*)+",
        "",
        raw,
        flags=re.IGNORECASE,
    )
    raw = re.split(AFFILIATION_RE, raw, maxsplit=1, flags=re.IGNORECASE)[0]
    raw = raw.strip(" |,;:+/-.")
    raw = re.sub(r"\s*-\s*", "-", raw)
    return clean_spaces(raw), starts_as_partner or explicit_partner_affiliation


def is_person_candidate(value: Any) -> bool:
    cleaned, is_partner = clean_person_candidate(value)
    if not cleaned or is_partner:
        return False
    tokens = ascii_key(cleaned).split()
    if len(tokens) not in {2, 3}:
        return False
    if any(token in NAME_STOP_WORDS or token.isdigit() for token in tokens):
        return False
    return all(len(token) >= 1 for token in tokens)


def build_inspector_candidates(raw_values: list[Any]) -> dict[str, str]:
    counts: Counter[str] = Counter()
    displays: defaultdict[str, Counter[str]] = defaultdict(Counter)
    for value in raw_values:
        for segment in split_raw_inspector_segments(value):
            cleaned, is_partner = clean_person_candidate(segment)
            tokens = ascii_key(cleaned).split()
            is_single_surname = (
                len(tokens) == 1
                and len(tokens[0]) >= 3
                and tokens[0] not in NAME_STOP_WORDS
            )
            if is_partner or (not is_single_surname and not is_person_candidate(cleaned)):
                continue
            key = ascii_key(cleaned)
            display = normalize_person_name(cleaned)
            counts[key] += 1
            displays[key][display] += 1

    candidates = {
        key: displays[key].most_common(1)[0][0]
        for key in counts
        if len(key.split()) > 1 or counts[key] >= 2
    }

    # Aceeasi persoana apare uneori cu prenumele inversate.
    by_token_set: defaultdict[tuple[str, ...], list[str]] = defaultdict(list)
    for key in candidates:
        tokens = key.split()
        if len(tokens) >= 2:
            by_token_set[tuple(sorted(tokens))].append(key)
    for keys in by_token_set.values():
        if len(keys) <= 1:
            continue
        canonical_key = max(keys, key=lambda item: counts[item])
        canonical_display = candidates[canonical_key]
        for key in keys:
            candidates[key] = canonical_display

    # Unifica doar erori ortografice evidente: aceeasi pozitie a numelui,
    # toate celelalte componente identice si o singura componenta foarte apropiata.
    canonical_keys: list[str] = []
    for key in sorted(candidates, key=lambda item: (-counts[item], item)):
        tokens = key.split()
        matched_key = None
        for canonical_key in canonical_keys:
            canonical_tokens = canonical_key.split()
            if len(tokens) != len(canonical_tokens) or len(tokens) < 2:
                continue
            different = [
                (left, right)
                for left, right in zip(tokens, canonical_tokens)
                if left != right
            ]
            if len(different) != 1:
                continue
            left, right = different[0]
            if min(len(left), len(right)) < 4:
                continue
            if SequenceMatcher(None, left, right).ratio() >= 0.80:
                matched_key = canonical_key
                break
        if matched_key:
            candidates[key] = candidates[matched_key]
        else:
            canonical_keys.append(key)

    # Un singur nume de familie este legat de forma completa doar daca
    # exista o unica persoana compatibila in registru.
    full_by_surname: defaultdict[str, list[str]] = defaultdict(list)
    for key in candidates:
        tokens = key.split()
        if len(tokens) >= 2:
            full_by_surname[tokens[0]].append(key)
    for key in list(candidates):
        if len(key.split()) != 1:
            continue
        matches = list(dict.fromkeys(full_by_surname.get(key, [])))
        if len(matches) == 1:
            candidates[key] = candidates[matches[0]]

    full_by_initial: defaultdict[tuple[str, str], list[str]] = defaultdict(list)
    for key in candidates:
        tokens = key.split()
        if len(tokens) >= 2 and len(tokens[1]) > 1:
            full_by_initial[(tokens[0], tokens[1][0])].append(key)
    for key in list(candidates):
        tokens = key.split()
        if len(tokens) == 2 and len(tokens[1]) == 1:
            matches = full_by_initial.get((tokens[0], tokens[1]), [])
            if len(matches) == 1:
                candidates[key] = candidates[matches[0]]
    return candidates


def match_inspector_candidates(
    value: Any,
    candidates: dict[str, str],
) -> tuple[list[str], bool]:
    cleaned, is_partner = clean_person_candidate(value)
    if not cleaned or is_partner:
        return [], False

    raw_key = ascii_key(cleaned)
    matches: list[tuple[int, int, str]] = []
    for key, display in candidates.items():
        for match in re.finditer(rf"(?<![a-z0-9]){re.escape(key)}(?![a-z0-9])", raw_key):
            matches.append((match.start(), match.end(), display))
    matches.sort(key=lambda item: (item[0], -(item[1] - item[0])))

    selected: list[tuple[int, int, str]] = []
    for start, end, display in matches:
        if any(not (end <= old_start or start >= old_end) for old_start, old_end, _ in selected):
            continue
        selected.append((start, end, display))
    selected.sort()

    if selected:
        names = list(dict.fromkeys(display for _, _, display in selected))
        covered = sum(end - start for start, end, _ in selected)
        return names, covered / max(1, len(raw_key)) < 0.6

    if is_person_candidate(cleaned):
        key = ascii_key(cleaned)
        return [candidates.get(key, normalize_person_name(cleaned))], False
    return [], bool(raw_key)


def split_inspectors(value: Any, candidates: dict[str, str]) -> tuple[list[dict[str, str]], bool]:
    if value in (None, ""):
        return [], False

    names: list[str] = []
    uncertain = False
    for segment in split_raw_inspector_segments(value):
        segment_names, segment_uncertain = match_inspector_candidates(segment, candidates)
        names.extend(segment_names)
        uncertain = uncertain or segment_uncertain

    deduplicated: dict[str, str] = {}
    for name in names:
        deduplicated.setdefault(ascii_key(name), name)
    return [{"nume": name} for name in deduplicated.values()], uncertain


def infer_entity_type(value: Any) -> str:
    key = ascii_key(value)
    if "ocol" in key and "silvic" in key:
        return "ocol silvic"
    if any(token in key for token in ("srl", "sa ", "sc ", "if ", "pfa", "societ")):
        return "operator economic"
    if "primar" in key or "uat" in key:
        return "primarie"
    if "fond cinegetic" in key or "avps" in key:
        return "gestionar fond cinegetic"
    return "persoana fizica"


def duplicate_key(row: ParsedRow) -> tuple[Any, ...]:
    payload = row.payload
    return (
        row.date_start.isoformat(),
        row.date_end.isoformat(),
        payload.get("garda"),
        ascii_key(payload.get("localitate")),
        ascii_key(payload.get("entitate_controlata")),
        clean_spaces(payload.get("cui")),
        ascii_key(payload.get("numar_act_control")),
    )


def read_source_rows(path: Path) -> tuple[list[tuple[int, tuple[Any, ...]]], list[Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if SOURCE_SHEET not in workbook.sheetnames:
        raise ValueError(f"Foaia obligatorie {SOURCE_SHEET!r} nu exista. Foi: {workbook.sheetnames}")
    sheet = workbook[SOURCE_SHEET]
    rows = []
    inspector_values = []
    for excel_row, values in enumerate(
        sheet.iter_rows(min_row=DATA_START_ROW, max_col=EXPECTED_COLUMNS, values_only=True),
        start=DATA_START_ROW,
    ):
        if not any(value not in (None, "") for value in values):
            continue
        values = tuple(values)
        rows.append((excel_row, values))
        inspector_values.append(values[10])
    workbook.close()
    return rows, inspector_values


def parse_rows(path: Path) -> tuple[list[ParsedRow], dict[str, Any]]:
    source_rows, inspector_values = read_source_rows(path)
    candidates = build_inspector_candidates(inspector_values)
    parsed_rows: list[ParsedRow] = []
    issues: defaultdict[str, list[Any]] = defaultdict(list)
    distributions = {
        "guards": Counter(),
        "months": Counter(),
        "types": Counter(),
        "results": Counter(),
        "inspectors": Counter(),
        "team_sizes": Counter(),
        "coordinate_status": Counter(),
        "date_status": Counter(),
    }

    for excel_row, values in source_rows:
        (
            nr, raw_date, raw_guard, raw_locality, raw_county, raw_lat, raw_lon,
            raw_type, raw_entity, raw_cui, raw_inspectors, raw_findings,
            raw_result, raw_measures, raw_act,
        ) = values
        date_start, date_end, date_status = parse_control_dates(raw_date)
        distributions["date_status"][date_status or "valid"] += 1
        if not date_start or not date_end:
            issues["invalid_dates"].append({"row": excel_row, "value": json_value(raw_date)})
            continue
        if date_start.year != 2026 or date_end.year != 2026:
            issues["dates_outside_2026"].append({
                "row": excel_row,
                "value": json_value(raw_date),
                "date_start": date_start.isoformat(),
                "date_end": date_end.isoformat(),
            })
        if date_start > date.today() or date_end > date.today():
            issues["future_dates"].append({
                "row": excel_row,
                "value": json_value(raw_date),
                "date_start": date_start.isoformat(),
                "date_end": date_end.isoformat(),
            })

        guard = normalize_guard(raw_guard)
        if not guard:
            issues["missing_or_unknown_guard"].append({"row": excel_row, "value": json_value(raw_guard)})

        gps, coordinate_status = normalize_coordinates(raw_lat, raw_lon)
        distributions["coordinate_status"][coordinate_status] += 1
        if not gps:
            issues["invalid_coordinates"].append({
                "row": excel_row,
                "lat": json_value(raw_lat),
                "lon": json_value(raw_lon),
                "reason": coordinate_status,
            })

        team, uncertain_team = split_inspectors(raw_inspectors, candidates)
        distributions["team_sizes"][str(len(team))] += 1
        if not team:
            issues["missing_inspectors"].append({"row": excel_row, "value": json_value(raw_inspectors)})
        elif uncertain_team:
            issues["uncertain_inspector_split"].append({"row": excel_row, "value": json_value(raw_inspectors)})

        control_type = normalize_control_type(raw_type)
        result, result_group = normalize_result(raw_result)
        cui = clean_spaces(raw_cui)
        if cui.endswith(".0"):
            cui = cui[:-2]

        original_values = {
            "nr_crt": json_value(nr),
            "data_controlului": json_value(raw_date),
            "garda_forestiera": json_value(raw_guard),
            "localitate": json_value(raw_locality),
            "judet": json_value(raw_county),
            "latitudine": json_value(raw_lat),
            "longitudine": json_value(raw_lon),
            "tip_control": json_value(raw_type),
            "entitate_controlata": json_value(raw_entity),
            "cui": json_value(raw_cui),
            "agenti_constatatori": json_value(raw_inspectors),
            "constatare": json_value(raw_findings),
            "rezultat": json_value(raw_result),
            "masuri_dispuse": json_value(raw_measures),
            "numar_act_control": json_value(raw_act),
        }
        payload = {
            "source_kind": "real_excel",
            "source_workbook": path.name,
            "source_sheet": SOURCE_SHEET,
            "source_row": excel_row,
            "source_number": clean_spaces(nr),
            "import_batch": IMPORT_BATCH,
            "source_original": original_values,
            "data_control": date_start.isoformat(),
            "date_start": date_start.isoformat(),
            "date_end": date_end.isoformat(),
            "data_control_original": json_value(raw_date),
            "garda": guard,
            "garda_original": json_value(raw_guard),
            "judet": normalize_judet(raw_county),
            "judet_original": json_value(raw_county),
            "localitate": normalize_display_text(raw_locality),
            "localitate_original": json_value(raw_locality),
            "reper": None,
            "gps": gps,
            "coordinate_status": coordinate_status,
            "control_type": control_type,
            "tip_control_original": json_value(raw_type),
            "entitate_controlata": normalize_display_text(raw_entity),
            "entitate_controlata_original": json_value(raw_entity),
            "tip_entitate": infer_entity_type(raw_entity),
            "cui": cui or None,
            "echipa": team,
            "agenti_constatatori_original": json_value(raw_inspectors),
            "constatari": clean_spaces(raw_findings) or None,
            "constatare_originala": json_value(raw_findings),
            "result": result,
            "result_original": json_value(raw_result),
            "result_group": result_group,
            "masuri_dispuse": clean_spaces(raw_measures) or None,
            "masuri_dispuse_original": json_value(raw_measures),
            "numar_act_control": clean_spaces(raw_act) or None,
            "documente": [clean_spaces(raw_act)] if clean_spaces(raw_act) else [],
            "financial_data_available": False,
            "cuantum_amenda_ron": None,
            "valoare_prejudiciu_ron": None,
        }

        parsed = ParsedRow(
            excel_row=excel_row,
            source_number=clean_spaces(nr),
            date_start=date_start,
            date_end=date_end,
            control_type=control_type,
            result=result,
            payload=payload,
        )
        parsed_rows.append(parsed)
        distributions["guards"][guard or "Necunoscuta"] += 1
        distributions["months"][date_start.strftime("%Y-%m")] += 1
        distributions["types"][control_type] += 1
        distributions["results"][result] += 1
        for member in team:
            distributions["inspectors"][member["nume"]] += 1

    duplicate_groups: defaultdict[tuple[Any, ...], list[int]] = defaultdict(list)
    for row in parsed_rows:
        duplicate_groups[duplicate_key(row)].append(row.excel_row)
    potential_duplicates = [rows for rows in duplicate_groups.values() if len(rows) > 1]

    report = {
        "source_file": str(path),
        "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "source_sheet": SOURCE_SHEET,
        "source_rows": len(source_rows),
        "valid_rows": len(parsed_rows),
        "ignored_rows": len(source_rows) - len(parsed_rows),
        "import_batch": IMPORT_BATCH,
        "unique_inspectors": len(distributions["inspectors"]),
        "inspector_participations": sum(distributions["inspectors"].values()),
        "unique_guards": len([key for key in distributions["guards"] if key != "Necunoscuta"]),
        "valid_coordinates": distributions["coordinate_status"]["valid"] + distributions["coordinate_status"]["swapped"],
        "potential_duplicate_groups": len(potential_duplicates),
        "potential_duplicate_rows": sum(len(rows) for rows in potential_duplicates),
        "distributions": {
            key: dict(counter.most_common())
            for key, counter in distributions.items()
        },
        "top_inspectors": distributions["inspectors"].most_common(20),
        "issues": {
            key: {"count": len(values), "samples": values[:50]}
            for key, values in issues.items()
        },
        "potential_duplicate_samples": potential_duplicates[:50],
    }
    return parsed_rows, report


def find_import_user(db) -> User:
    user = db.execute(
        select(User)
        .where(User.role.in_(["inspector_general", "admin"]))
        .order_by(User.id)
    ).scalars().first()
    if not user:
        raise RuntimeError("Nu exista user cu rol inspector_general sau admin. Importul a fost oprit.")
    return user


def write_report(report: dict[str, Any]) -> None:
    REPORT_FILE.parent.mkdir(parents=True, exist_ok=True)
    REPORT_FILE.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def import_controls(rows: list[ParsedRow], report: dict[str, Any]) -> None:
    if not IMPORT_CONFIRM:
        print("DRY RUN: datele NU au fost modificate. Pentru import seteaza GFN_REAL_IMPORT_CONFIRM=true.")
        return
    if not IMPORT_REPLACE:
        raise RuntimeError(
            "Importul real necesita GFN_REAL_IMPORT_REPLACE=true pentru inlocuirea explicita a datelor demo."
        )

    with SessionLocal() as db:
        user = find_import_user(db)
        old_total = db.scalar(select(func.count()).select_from(Control)) or 0
        try:
            db.execute(delete(ControlAuditLog))
            deleted_controls = db.execute(delete(Control)).rowcount or 0
            controls = []
            for index, row in enumerate(rows):
                created_at = datetime.combine(row.date_start, time(12, 0, 0)).replace(microsecond=index % 1_000_000)
                controls.append(
                    Control(
                        created_at=created_at,
                        created_by_user_id=user.id,
                        result=row.result,
                        control_type=row.control_type,
                        payload=row.payload,
                    )
                )
            db.add_all(controls)
            db.commit()
        except Exception:
            db.rollback()
            raise

        report["database"] = {
            "controls_before_import": old_total,
            "deleted_existing_controls": deleted_controls,
            "inserted_controls": len(controls),
            "created_by_user_id": user.id,
            "created_by_role": user.role,
        }
        write_report(report)


def main() -> None:
    if not IMPORT_FILE.exists():
        raise FileNotFoundError(f"Nu gasesc fisierul Excel: {IMPORT_FILE}")
    rows, report = parse_rows(IMPORT_FILE)
    write_report(report)
    print(json.dumps({
        "source_rows": report["source_rows"],
        "valid_rows": report["valid_rows"],
        "ignored_rows": report["ignored_rows"],
        "unique_inspectors": report["unique_inspectors"],
        "inspector_participations": report["inspector_participations"],
        "team_sizes": report["distributions"]["team_sizes"],
        "unique_guards": report["unique_guards"],
        "valid_coordinates": report["valid_coordinates"],
        "coordinate_status": report["distributions"]["coordinate_status"],
        "guards": report["distributions"]["guards"],
        "months": report["distributions"]["months"],
        "top_inspectors": report["top_inspectors"][:10],
        "issues": {key: item["count"] for key, item in report["issues"].items()},
        "potential_duplicate_groups": report["potential_duplicate_groups"],
        "report_file": str(REPORT_FILE),
    }, ensure_ascii=False, indent=2))
    import_controls(rows, report)
    if IMPORT_CONFIRM:
        print(f"Import finalizat: {len(rows)} controale reale.")


if __name__ == "__main__":
    main()
